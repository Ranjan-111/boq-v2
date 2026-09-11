"""T101 XLSX writer tests — the csv_export contract, openpyxl-rendered.

Pinned:
  * byte-determinism: two calls with the same rows → identical bytes (the
    manifest sha256 is meaningful only if this holds; openpyxl's default
    save path stamps now() into docProps + zip entry times — the writer
    normalizes both),
  * the shared validate_export gate: no approval, stale digest, bad totals,
    mixed currencies all refuse exactly like the CSV writer,
  * grand total = exact integer-minor sum, formatted 2dp via Decimal,
  * the workbook loads back via openpyxl with the expected cell values
    (header, items, 6dp quantity, 2dp money strings, markup int),
  * quantity/money strings never show float artifacts.
"""
from __future__ import annotations

import io
import time
from dataclasses import dataclass, replace
from decimal import Decimal

import openpyxl  # type: ignore[import-untyped]  # no py.typed yet
import pytest

from core.domain.enums import BoqStatus
from exports.csv_export import ExportApproval, rows_digest
from exports.xlsx_export import xlsx_bytes

UUID_A = "11111111-1111-1111-1111-111111111111"
UUID_B = "22222222-2222-2222-2222-222222222222"


@dataclass(frozen=True, slots=True)
class Row:
    """The CsvExportRow shape (a local fixture — exports is Protocol-typed)."""

    code: str = "2.1.1"
    description: str = "Brick wall 230mm thick"
    unit: str = "m"
    quantity: Decimal = Decimal("12.5")
    rate_minor: int = 85_000
    markup_bp: int = 750
    total_minor: int = 1_142_188  # 12.5 x 850.00 + 7.5% markup (banker's)
    currency: str = "INR"
    measurement_ids: tuple[str, ...] = (UUID_A,)


def approved(rows: list[Row]) -> ExportApproval:
    return ExportApproval("boq-1", "approval-1", rows_digest(rows),
                          BoqStatus.APPROVED)


def two_rows() -> list[Row]:
    return [
        Row(),
        Row(quantity=Decimal("4.75"), total_minor=434_031,
            measurement_ids=(UUID_B,)),
    ]


class TestDeterminism:
    def test_two_calls_are_byte_identical(self) -> None:
        rows = two_rows()
        approval = approved(rows)
        first = xlsx_bytes(rows, approval=approval)
        time.sleep(1.1)  # openpyxl's default path would now() here
        assert xlsx_bytes(rows, approval=approval) == first

    def test_reexport_with_exported_status_is_identical(self) -> None:
        rows = [Row()]
        approval = approved(rows)
        assert xlsx_bytes(rows, approval=approval) == xlsx_bytes(
            rows, approval=replace(approval, status=BoqStatus.EXPORTED))


class TestValidationGate:
    """The writers share validate_export — the CSV refusal set, pinned here."""

    def test_no_approval_refused(self) -> None:
        with pytest.raises(ValueError, match="approval"):
            xlsx_bytes(two_rows())

    @pytest.mark.parametrize("status", [BoqStatus.DRAFT, BoqStatus.IN_REVIEW,
                                        BoqStatus.REVIEWED,
                                        BoqStatus.STALE_APPROVED])
    def test_unapproved_state_refused(self, status: BoqStatus) -> None:
        rows = two_rows()
        with pytest.raises(ValueError, match="status"):
            xlsx_bytes(rows, approval=replace(approved(rows), status=status))

    def test_stale_digest_refused(self) -> None:
        rows = two_rows()
        changed = [replace(rows[0], description="changed")]
        with pytest.raises(ValueError, match="stale"):
            xlsx_bytes(changed, approval=approved(rows))

    def test_bad_total_refused(self) -> None:
        rows = [replace(Row(), total_minor=1)]  # not recomputable
        with pytest.raises(ValueError, match="recomputable"):
            xlsx_bytes(rows, approval=approved(rows))

    def test_mixed_currency_refused(self) -> None:
        rows = [Row(), Row(currency="USD",
                           measurement_ids=(UUID_B,))]
        with pytest.raises(ValueError, match="mixed currencies"):
            xlsx_bytes(rows, approval=approved(rows))

    def test_empty_rows_refused(self) -> None:
        with pytest.raises(ValueError, match="no priced rows"):
            xlsx_bytes([], approval=ExportApproval(
                "boq-1", "approval-1", rows_digest([]), BoqStatus.APPROVED))


class TestWorkbookContent:
    def _sheet(self, data: bytes) -> openpyxl.Worksheet:
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert wb.active is not None
        return wb.active

    def test_loads_back_with_expected_cells(self) -> None:
        data = xlsx_bytes(two_rows(), approval=approved(two_rows()))
        ws = self._sheet(data)
        assert ws["A2"].value == "sr_no"
        assert ws["J2"].value == "measurement_refs"
        assert ws["A3"].value == 1 and ws["A4"].value == 2
        assert ws["B3"].value == "2.1.1"
        assert ws["E3"].value == "12.500000"  # 6dp quantity string
        assert ws["F3"].value == "850.00"  # 2dp rate string
        assert ws["G3"].value == 750  # markup: plain int
        assert ws["H3"].value == "11421.88"  # 2dp total string
        assert ws["I3"].value == "INR"
        assert ws["J3"].value == UUID_A

    def test_grand_total_row_is_exact_integer_sum(self) -> None:
        rows = two_rows()
        data = xlsx_bytes(rows, approval=approved(rows))
        ws = self._sheet(data)
        last = ws.max_row
        assert ws.cell(row=last, column=2).value == "GRAND TOTAL"
        expected_minor = rows[0].total_minor + rows[1].total_minor
        assert ws.cell(row=last, column=8).value == str(
            Decimal(expected_minor) / Decimal(100))
        assert (ws.cell(row=last, column=8).value
                == f"{expected_minor // 100}.{expected_minor % 100:02d}")

    def test_title_frozen_header_and_bold_rows(self) -> None:
        rows = two_rows()
        data = xlsx_bytes(rows, approval=approved(rows),
                           title="Acme BOQ")
        ws = self._sheet(data)
        assert ws["A1"].value == "Acme BOQ"
        assert ws.freeze_panes == "A3"  # title + header stay visible
        assert ws["A1"].font.bold is True
        assert ws["A2"].font.bold is True
        assert ws.cell(row=ws.max_row, column=2).font.bold is True
        # Column widths set (no default None collapse).
        assert ws.column_dimensions["C"].width == 48

    def test_no_float_artifacts_in_money(self) -> None:
        # 0.07 m x 850.00 = 59.50 base; markup 4.46 -> 63.96 exact —
        # a float path would show 63.96000000000001-style strings.
        rows = [Row(quantity=Decimal("0.07"), total_minor=6_396)]
        data = xlsx_bytes(rows, approval=approved(rows))
        ws = self._sheet(data)
        assert ws["H3"].value == "63.96"
        assert ws["E3"].value == "0.070000"
