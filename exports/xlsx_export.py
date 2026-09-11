"""XLSX export (T101) — priced BOQ rows → deterministic openpyxl workbook bytes.

Same contract as csv_export: same Protocol rows, same validate_export gate,
same Decimal/int-only money formatting. The xlsx extra (openpyxl, MIT) is
imported here only — never re-exported.

Determinism contract: openpyxl's Workbook.save() stamps properties.modified
with now() and the zipfile layer stamps every entry with the wall clock —
neither can be disabled via its public API. This writer therefore
(1) builds through ExcelWriter directly (skipping save_workbook's now()
stamp) with fixed DocumentProperties, then
(2) rewrites the zip entries in the same order with a FIXED ZipInfo
timestamp, so two calls with the same rows are byte-identical and the
manifest sha256 stays meaningful.
"""
from __future__ import annotations

import io
import zipfile
from collections.abc import Sequence
from datetime import UTC, datetime

# openpyxl ships no py.typed/stubs package yet — per-line ignores, the
# ingestion/dxf (ezdxf) idiom. Swap for a types- package when one is blessed.
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.packaging.core import DocumentProperties  # type: ignore[import-untyped]
from openpyxl.styles import Font  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.writer.excel import ExcelWriter  # type: ignore[import-untyped]

from exports.csv_export import (
    HEADER,
    CsvExportRow,
    ExportApproval,
    fmt_money_minor,
    fmt_quantity,
    validate_export,
)

# Fixed stamp so sha256(xlsx_bytes) is stable across runs (module docstring).
_FIXED_STAMP = datetime(2026, 1, 1, tzinfo=UTC)
_BOLD = Font(bold=True)
_GRAND_TOTAL = "GRAND TOTAL"

# Column widths (chars) — generous for UUID-laden measurement_refs.
_COL_WIDTHS = (8, 12, 48, 8, 14, 12, 11, 14, 10, 44)


def _normalize_zip(data: bytes) -> bytes:
    """Rewrite every zip entry with a fixed timestamp (same order, same bytes)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename,
                                   date_time=(2026, 1, 1, 0, 0, 0))
            fixed.compress_type = info.compress_type
            fixed.external_attr = info.external_attr
            fixed.internal_attr = info.internal_attr
            fixed.create_system = info.create_system
            dst.writestr(fixed, src.read(info.filename))
    return out.getvalue()


def xlsx_bytes(
    rows: Sequence[CsvExportRow], *, approval: ExportApproval | None = None,
    title: str = "Bill of Quantities",
) -> bytes:
    """Render BOQ rows as xlsx bytes (sha256-able, byte-deterministic).

    Columns mirror the CSV exactly — quantity as a 6dp string, rate/total as
    2dp strings via Decimal division (never float cells, never float math).
    V1 rows are flat (no section grouping in the service's export adapter), so
    a single GRAND TOTAL row follows the items. Sheet layout: title (1),
    header (2, frozen), items, grand total.
    """
    validate_export(rows, approval)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "BOQ"
    ws.append([title])
    ws.append(HEADER)
    total_minor = 0
    for i, row in enumerate(rows, start=1):
        ws.append([
            i,
            row.code,
            row.description,
            row.unit,
            fmt_quantity(row.quantity),
            fmt_money_minor(row.rate_minor),
            row.markup_bp,
            fmt_money_minor(row.total_minor),
            row.currency,
            ";".join(row.measurement_ids),
        ])
        total_minor += row.total_minor
    grand_row = 2 + len(rows) + 1  # title + header + items + 1
    ws.append(["", _GRAND_TOTAL, "", "", "", "", "",
               fmt_money_minor(total_minor), rows[0].currency, ""])
    for sheet_row in (1, 2, grand_row):  # title, header, grand total
        for cell in ws[sheet_row]:
            cell.font = _BOLD
    ws.freeze_panes = "A3"  # title + header stay visible while scrolling
    for i, width in enumerate(_COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    wb.properties = DocumentProperties(
        creator="", title=title,
        created=_FIXED_STAMP, modified=_FIXED_STAMP, lastModifiedBy="")

    # The deterministic save path: ExcelWriter (no save_workbook now() stamp)
    # into a fresh zip, then the fixed-timestamp rewrite.
    raw = io.BytesIO()
    archive = zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED, allowZip64=True)
    ExcelWriter(wb, archive).save()
    return _normalize_zip(raw.getvalue())
